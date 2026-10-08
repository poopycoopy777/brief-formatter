"""Tests for the JDF 1987 brief formatter.

Runs on the standard library alone, so either interpreter works:

    python -m unittest tests.test_jdf_brief -v
    python tests/test_jdf_brief.py
"""

from __future__ import annotations

import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from jdf_brief.build import _split_label, build_brief  # noqa: E402
from jdf_brief.citations import italic_segments  # noqa: E402
from jdf_brief.parse import (  # noqa: E402
    BODY,
    ISSUE,
    LIST,
    MAJOR,
    SUB,
    parse,
)


def _italic_text(text: str) -> list[str]:
    """Return just the italicised chunks of *text*."""
    return [chunk for chunk, is_italic in italic_segments(text) if is_italic]


# ---------------------------------------------------------------------------
# Citation italics
# ---------------------------------------------------------------------------

class TestCitationItalics(unittest.TestCase):
    def test_case_names_and_reporters_are_italicised(self):
        cases = [
            (
                "Corder v. Folds, 292 P.3d 1177, 1181 (Colo. App. 2012).",
                ["Corder v. Folds", "292 P.3d 1177"],
            ),
            (
                "Betterview Investments, LLC v. Public Service Co., 198 P.3d 1258, 1262 (Colo. App. 2008).",
                ["Betterview Investments, LLC v. Public Service Co.", "198 P.3d 1258"],
            ),
            (
                "Ion Media Networks, Inc. v. West, 2025 COA 66, P 13-14.",
                ["Ion Media Networks, Inc. v. West", "2025 COA 66"],
            ),
            (
                "Norton v. Rocky Mountain Planned Parenthood, Inc., 2018 CO 3, P 7, 409 P.3d 331, 334.",
                [
                    "Norton v. Rocky Mountain Planned Parenthood, Inc.",
                    "2018 CO 3",
                    "409 P.3d 331",
                ],
            ),
            (
                "See Dunlap v. Colo. Springs Cablevision, Inc., 829 P.2d 1286, 1290 (Colo. 1992).",
                ["Dunlap v. Colo. Springs Cablevision, Inc.", "829 P.2d 1286"],
            ),
            (
                "Finnie v. Jefferson Cnty. Sch. Dist. R-1, 79 P.3d 1253, 1259 (Colo. 2003).",
                ["Finnie v. Jefferson Cnty. Sch. Dist. R-1", "79 P.3d 1253"],
            ),
            (
                "SMB Advertising, Inc. v. City of Boulder, 2026 COA 25M, P 21.",
                ["SMB Advertising, Inc. v. City of Boulder", "2026 COA 25M"],
            ),
        ]
        for text, expected in cases:
            with self.subTest(text=text[:48]):
                self.assertEqual(_italic_text(text), expected)

    def test_ordinary_prose_is_left_roman(self):
        cases = [
            "The court entered judgment on May 22, 2026. CF, p 123-131.",
            "Under C.R.S. 24-31-902(2)(a), the agency shall release the recordings.",
            "WHEREFORE, Appellant respectfully requests that this Court:",
            "The file is named my_brief_final.txt.",
            "Appellant filed his Complaint on December 8, 2025.",
        ]
        for text in cases:
            with self.subTest(text=text[:48]):
                self.assertEqual(_italic_text(text), [])

    def test_explicit_italic_markup(self):
        self.assertEqual(_italic_text("The term *de novo* applies."), ["de novo"])
        self.assertEqual(_italic_text("The term _de novo_ applies."), ["de novo"])

    def test_markup_is_consumed_and_auto_detection_continues(self):
        chunks = italic_segments("*Ion Media Networks, Inc. v. West*, 2025 COA 66.")
        self.assertIn(("Ion Media Networks, Inc. v. West", True), chunks)
        self.assertTrue(any(c == "2025 COA 66" and i for c, i in chunks))
        self.assertNotIn("*", "".join(c for c, _ in chunks))


# ---------------------------------------------------------------------------
# Paragraph splitting
# ---------------------------------------------------------------------------

class TestLabelSplitting(unittest.TestCase):
    def test_flattened_label_is_split_out(self):
        text = (
            "I certify that this brief complies with C.A.R. 28 and 32. Including: "
            "Word Limits:         My brief has 4,046 words."
        )
        # "Including:" is a clause lead-in, not a section label, so it stays
        # with the sentence it belongs to.
        self.assertEqual(
            _split_label(text),
            [
                "I certify that this brief complies with C.A.R. 28 and 32. Including:",
                "Word Limits:",
                "My brief has 4,046 words.",
            ],
        )

    def test_colon_inside_a_sentence_is_not_a_label(self):
        text = (
            'The operative holding rests on the CCJRA: "All other audio redactions '
            'are in compliance." CF, p 118.'
        )
        self.assertEqual(_split_label(text), [text])

    def test_sentence_initial_word_is_not_a_label(self):
        text = "Appellant filed his complaint. The district court dismissed it."
        self.assertEqual(_split_label(text), [text])

    def test_label_at_start_of_paragraph_is_split(self):
        self.assertEqual(
            _split_label("Standard of Review:   Reviewed de novo."),
            ["Standard of Review:", "Reviewed de novo."],
        )

    def test_abbreviation_does_not_create_a_false_boundary(self):
        text = "Under C.R.S. 24-31-902(2)(a), the agency shall release the recordings."
        self.assertEqual(_split_label(text), [text])


# ---------------------------------------------------------------------------
# Structure detection
# ---------------------------------------------------------------------------

CAPTION = """COURT OF APPEALS, STATE OF COLORADO
2 East 14th Avenue, Denver, Colorado 80203
District Court Case Number: 2025CV111
Plaintiff-Appellant:
HARRY COOPER II
v.
Defendants-Appellees:
JENNA INGO
\tCourt of Appeals Case Number: 2026CA1239
\tAPPELLANT'S OPENING BRIEF
"""


class TestStructureDetection(unittest.TestCase):
    def test_caption_and_toc_are_skipped(self):
        text = CAPTION + """
\t1. Certificate of Compliance
I certify the brief complies.

Table of Content
Table of Authorities: Pg. 3
Issues on Appeal: Pg. 4

Table of Authorities
Cases
Statutes
Court Rules
Other Authorities Cited

3. ISSUES PRESENTED FOR REVIEW
1. Whether the district court erred.
2. Whether the dismissal must be reversed.
"""
        result = parse(text)
        self.assertTrue(result.skipped_caption)
        self.assertTrue(result.skipped_toc)

        headings = [b.text for b in result.blocks if b.kind == MAJOR]
        self.assertEqual(
            headings, ["Certificate of Compliance", "ISSUES PRESENTED FOR REVIEW"]
        )

        joined = " ".join(b.text for b in result.blocks)
        for residue in ("Table of Author", "Court Rules", "Pg. 9", "Statutes"):
            self.assertNotIn(residue, joined)

    def test_issue_headings_are_their_own_block_kind(self):
        text = """1. ARGUMENT
ISSUE 1: THE COURT APPLIED THE WRONG STATUTE.
Standard of Review
Reviewed de novo.
ISSUE 2 - THE DISMISSAL CANNOT STAND.
"""
        result = parse(text)
        issues = [b.text for b in result.blocks if b.kind == ISSUE]
        self.assertEqual(
            issues,
            ["THE COURT APPLIED THE WRONG STATUTE.", "THE DISMISSAL CANNOT STAND."],
        )

    def test_issues_presented_is_a_major_heading_not_an_issue(self):
        text = "2. ISSUES PRESENTED FOR REVIEW\nSomething follows.\n"
        result = parse(text)
        majors = [b.text for b in result.blocks if b.kind == MAJOR]
        self.assertEqual(majors, ["ISSUES PRESENTED FOR REVIEW"])
        self.assertEqual([b for b in result.blocks if b.kind == ISSUE], [])

    def test_lettered_subsections_keep_their_letter(self):
        text = """1. STATEMENT OF THE CASE
A. Jurisdiction
The court entered judgment.
B. Nature of the Case
This is a records action.
"""
        result = parse(text)
        subs = [(b.letter, b.text) for b in result.blocks if b.kind == SUB]
        self.assertEqual(subs, [("A.", "Jurisdiction"), ("B.", "Nature of the Case")])

    def test_one_paragraph_per_line_is_not_merged(self):
        # A brief written with each paragraph on its own source line must not
        # have those paragraphs joined together.
        text = """1. STATEMENT OF THE CASE
The district court entered judgment on May 22, 2026. CF, p 123.
Appellant timely filed his Notice of Appeal on June 24, 2026. CF, p 142-146.
The April 5 order was interlocutory and contemplated further proceedings.
"""
        result = parse(text)
        bodies = [b.text for b in result.blocks if b.kind == BODY]
        self.assertEqual(len(bodies), 3)
        self.assertTrue(bodies[0].startswith("The district court entered judgment"))
        self.assertTrue(bodies[1].startswith("Appellant timely filed"))
        self.assertTrue(bodies[2].startswith("The April 5 order"))

    def test_hard_wrapped_prose_is_still_joined(self):
        # Mid-sentence line breaks (as a PDF text extraction produces) must
        # still be re-joined into one paragraph.
        text = """1. STATEMENT OF THE CASE
The district court entered its final order dismissing all four claims on May
22, 2026. CF, p 123-131. Appellant timely filed his Notice of Appeal on June
24, 2026. CF, p 142-146.
"""
        result = parse(text)
        bodies = [b.text for b in result.blocks if b.kind == BODY]
        self.assertEqual(len(bodies), 1)
        self.assertIn("May 22, 2026", bodies[0])
        self.assertIn("June 24, 2026", bodies[0])

    def test_duplicate_consecutive_paragraph_is_collapsed(self):
        em = "\u2014"
        first = (
            "The sole privacy exception, \u00a7 24-31-902(2)(b)(II)(A), decides this "
            "appeal on three features. First, it is written in video and only in video: "
            "subsection (2)(a) uses both nouns " + em + " \"video and audio\" " + em + " "
            "while (2)(b)(II)(A) reaches any video that raises substantial privacy "
            "concerns and provides that it does not permit the removal of any portion "
            "of the video that was recorded."
        )
        second = first.replace(em, "-") + " Extra trailing sentence."
        text = "1. ARGUMENT\n%s\n%s\n" % (first, second)

        result = parse(text)
        bodies = [b.text for b in result.blocks if b.kind == BODY]
        self.assertEqual(len(bodies), 1)
        self.assertEqual(result.removed_duplicates, 1)
        # The longer copy is kept.
        self.assertTrue(bodies[0].endswith("Extra trailing sentence."))

    def test_distant_repeated_sentence_is_kept(self):
        # The same sentence legitimately appears in the summary and again in
        # the argument; only consecutive duplicates are collapsed.
        opener = (
            "The Integrity Act commands release of all unedited video and audio "
            "recordings of the incident under the governing statute."
        )
        filler = " ".join(["Filler sentence about the record on appeal."] * 8)
        text = "1. SUMMARY\n%s\n%s\n2. ARGUMENT\n%s\n%s\n" % (
            opener, filler, opener, filler,
        )
        result = parse(text)
        self.assertEqual(result.removed_duplicates, 0)
        joined = " ".join(b.text for b in result.blocks if b.kind == BODY)
        self.assertEqual(joined.count("commands release of all unedited"), 2)

    def test_collapse_can_be_disabled(self):
        body = (
            "A paragraph long enough to qualify for duplicate detection, repeated "
            "here so that the comparison has enough words to work with comfortably "
            "and reliably. It runs to more than twenty five words in total overall."
        )
        text = "1. ARGUMENT\n%s\n%s\n" % (body, body)
        self.assertEqual(parse(text).removed_duplicates, 1)
        self.assertEqual(
            parse(text, collapse_duplicates=False).removed_duplicates, 0
        )

    def test_numbered_issue_list_stays_a_list(self):
        text = """1. ISSUES PRESENTED FOR REVIEW
1. Whether the district court erred as a matter of law by sustaining continued muting of body-worn camera audio under the general discretionary withholding provisions of the CCJRA.
2. Whether the final order dismissing all four claims must be reversed.
3. CONCLUSION
The judgment should be reversed.
"""
        result = parse(text)
        lists = [b.text for b in result.blocks if b.kind == LIST]
        self.assertEqual(len(lists), 2)
        self.assertTrue(lists[0].startswith("1. Whether"))
        self.assertTrue(lists[1].startswith("2. Whether"))
        self.assertTrue(
            any(b.kind == MAJOR and b.text == "CONCLUSION" for b in result.blocks)
        )

    def test_explicit_markers_override_detection(self):
        text = "# My Section\n## My Issue\n### A. My Sub\ntext\n> quoted text\n"
        result = parse(text)
        kinds = [b.kind for b in result.blocks]
        self.assertEqual(kinds[0], MAJOR)
        self.assertEqual(kinds[1], ISSUE)
        self.assertEqual(kinds[2], SUB)


# ---------------------------------------------------------------------------
# DOCX output
# ---------------------------------------------------------------------------

BRIEF = """1. Certificate of Compliance
I certify that this brief complies with C.A.R. 28 and 32.[^ The limit is 9,500 words.]

2. STATEMENT OF THE CASE
A. Jurisdiction
The district court entered judgment on May 22, 2026. CF, p 123. The court relied on Ion Media Networks, Inc. v. West, 2025 COA 66, P 36.

:: This indented paragraph should sit at a half inch with no first-line indent.

3. CONCLUSION
The judgment should be reversed.
"""


class TestDocxOutput(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:
            import docx  # noqa: F401
        except ImportError:  # pragma: no cover - depends on the interpreter
            raise unittest.SkipTest(
                "python-docx is not installed in this interpreter; the DOCX "
                "tests need it, the parser and citation tests do not."
            )
        cls.tmp = Path(tempfile.mkdtemp(prefix="jdf-test-"))
        cls.docx = cls.tmp / "brief.docx"
        build_brief(BRIEF, cls.docx)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def _paragraphs(self):
        import docx

        return docx.Document(str(self.docx)).paragraphs

    def test_package_is_valid_and_has_footnotes(self):
        with zipfile.ZipFile(self.docx) as archive:
            self.assertIsNone(archive.testzip())
            names = archive.namelist()
            self.assertIn("word/document.xml", names)
            self.assertIn("word/footnotes.xml", names)
            self.assertTrue(any(n.startswith("word/footer") for n in names))

    def test_footnote_part_is_registered(self):
        with zipfile.ZipFile(self.docx) as archive:
            self.assertIn("footnotes.xml", archive.read("[Content_Types].xml").decode())
            self.assertIn(
                "footnotes", archive.read("word/_rels/document.xml.rels").decode()
            )
            self.assertIn(
                "w:footnoteReference", archive.read("word/document.xml").decode()
            )
            footnotes = archive.read("word/footnotes.xml").decode()
            self.assertIn("The limit is 9,500 words.", footnotes)
            # One real footnote plus the separator and continuationSeparator.
            self.assertEqual(footnotes.count("<w:footnote "), 3)

    def test_page_setup_is_letter_with_one_inch_margins(self):
        import docx

        section = docx.Document(str(self.docx)).sections[0]
        self.assertAlmostEqual(section.page_width.inches, 8.5, places=2)
        self.assertAlmostEqual(section.page_height.inches, 11.0, places=2)
        for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            self.assertAlmostEqual(getattr(section, attr).inches, 1.0, places=2)

    def test_typography_matches_the_jdf_sample(self):
        from docx.shared import Pt

        paragraphs = self._paragraphs()

        major = paragraphs[0]
        self.assertTrue(major.text.startswith("1.\tCertificate of Compliance"))
        self.assertEqual(major.runs[0].font.name, "Garamond")
        self.assertEqual(major.runs[0].font.size, Pt(14))
        self.assertTrue(major.runs[0].font.bold)
        self.assertAlmostEqual(major.paragraph_format.left_indent.inches, 0.0, places=2)

        body = paragraphs[1]
        self.assertAlmostEqual(
            body.paragraph_format.first_line_indent.inches, 0.5, places=2
        )
        self.assertAlmostEqual(body.paragraph_format.line_spacing, 2.0, places=2)
        self.assertEqual(body.runs[0].font.size, Pt(12))

    def test_sections_are_renumbered_sequentially(self):
        paragraphs = self._paragraphs()
        majors = []
        for paragraph in paragraphs:
            if not paragraph.runs:
                continue
            run = paragraph.runs[0]
            if run.font.size and run.font.size.pt == 14 and run.font.bold:
                majors.append(paragraph.text)
        self.assertTrue(majors[0].startswith("1.\tCertificate of Compliance"))
        self.assertTrue(majors[1].startswith("2.\tSTATEMENT OF THE CASE"))
        self.assertTrue(majors[2].startswith("3.\tCONCLUSION"))

    def test_legal_indent_has_no_first_line(self):
        paragraphs = self._paragraphs()
        indented = [
            p for p in paragraphs
            if p.paragraph_format.left_indent is not None
            and abs(p.paragraph_format.left_indent.inches - 0.5) < 0.01
            and p.text.startswith("This indented paragraph")
        ]
        self.assertEqual(len(indented), 1)
        self.assertAlmostEqual(
            indented[0].paragraph_format.first_line_indent.inches, 0.0, places=2
        )

    def test_footer_carries_the_jdf_identifier(self):
        import docx

        footer = docx.Document(str(self.docx)).sections[0].footer
        text = "\n".join(p.text for p in footer.paragraphs)
        self.assertIn("JDF 1987", text)
        self.assertIn("Sample Opening Brief", text)
        # The page number is a live field, not literal text.
        self.assertNotIn("Page", text)


if __name__ == "__main__":
    unittest.main(verbosity=2)
